"""Check the persisted workbook and the stated cross-tool KPI contracts."""
import hashlib
import re
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from contract import DAYS, read_tables
from export_excel import export, KPI_LABELS, SHEETS, sla_rows
from verified_metrics import summarize, ticket_metrics


class OutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = read_tables("processed")
        cls.kpis = summarize(cls.data)
        paths = sorted(p for p in (ROOT / "data").rglob('*') if p.is_file())
        cls.before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
        cls.temp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.temp.name) / "analysis.xlsx"
        cls.result = export(cls.path)
        cls.workbook = load_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close()
        cls.temp.cleanup()

    def rows(self, sheet, index=0):
        ws = self.workbook[sheet]
        table = list(ws.tables.values())[index]
        first_col, first_row, last_col, last_row = range_boundaries(table.ref)
        headers = [ws.cell(first_row, column).value for column in range(first_col, last_col + 1)]
        return [dict(zip(headers, [ws.cell(row, column).value for column in range(first_col, last_col + 1)]))
                for row in range(first_row + 1, last_row + 1)]

    def test_workbook_is_real_complete_and_readable(self):
        self.assertTrue(self.path.is_file())
        self.assertEqual(self.workbook.sheetnames, SHEETS)
        self.assertEqual(self.result["charts"], 4)
        for ws in self.workbook:
            self.assertTrue(ws.freeze_panes)
            self.assertTrue(ws.tables)
            self.assertFalse(ws.merged_cells.ranges)
            self.assertGreater(len(self.rows(ws.title)), 0)
            for row in ws:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertNotIn(str(cell.value).strip().lower(), {"nan", "inf", "infinity", "-inf", "-infinity"})

    def test_every_excel_kpi_matches_python_and_saved_evidence(self):
        observed = {row["Metric"]: row["Value"] for row in self.rows("Executive_KPIs")}
        saved = pd.read_csv(ROOT / "data/analytics/verified_kpis.csv").set_index("metric").value
        self.assertEqual(len(observed), len(self.kpis))
        for key, value in self.kpis.items():
            self.assertTrue(np.isclose(observed[KPI_LABELS[key]], value, rtol=1e-12, atol=1e-9), key)
            self.assertTrue(np.isclose(observed[KPI_LABELS[key]], saved[key], rtol=1e-12, atol=1e-9), key)
        ws = self.workbook["Executive_KPIs"]
        for row in range(5, ws.max_row + 1):
            if "%" in ws.cell(row, 1).value:
                self.assertEqual(ws.cell(row, 2).number_format, "0.00%")
                self.assertGreaterEqual(ws.cell(row, 2).value, 0)
                self.assertLessEqual(ws.cell(row, 2).value, 1)

    def test_demand_sections_reconcile_and_first_change_is_blank(self):
        for index in range(5):
            rows = self.rows("Demand_Analysis", index)
            self.assertEqual(sum(row["Ticket_Count"] for row in rows), self.kpis["total_tickets"])
        monthly = self.rows("Demand_Analysis")
        self.assertIsNone(monthly[0]["Previous_Period"])
        self.assertIsNone(monthly[0]["Previous_Ticket_Count"])
        self.assertIsNone(monthly[0]["Change_Percent"])
        for before, after in zip(monthly, monthly[1:]):
            self.assertEqual(after["Previous_Ticket_Count"], before["Ticket_Count"])
            self.assertAlmostEqual(after["Change_Percent"], (after["Ticket_Count"] - before["Ticket_Count"]) / before["Ticket_Count"])
        weekday = self.rows("Demand_Analysis", 1)
        self.assertEqual(sum(row["Calendar_Days"] for row in weekday), len(DAYS))
        self.assertEqual(len(self.rows("Demand_Analysis", 2)), 24)

    def test_sla_eligibility_and_zero_eligible_cohort(self):
        rows = self.rows("SLA_Analysis")
        for row in rows:
            self.assertEqual(row["Eligible"], row["Met"] + row["Breached"])
            if row["Eligible"]:
                self.assertAlmostEqual(row["Compliance_Rate"], row["Met"] / row["Eligible"])
                self.assertAlmostEqual(row["Breach_Rate"], row["Breached"] / row["Eligible"])
        t = ticket_metrics(self.data["tickets"], self.data["sla_policies"])
        t["owner_team"] = "test_cohort"
        pending = t.loc[t.fr_sla_outcome.eq("PENDING")]
        self.assertGreater(len(pending), 0)
        fr = sla_rows(pending).query("SLA_Type == 'First Response'")
        self.assertTrue(fr.Eligible.eq(0).all())
        self.assertTrue(fr.Compliance_Rate.isna().all())
        self.assertTrue(fr.Breach_Rate.isna().all())

    def test_backlog_buckets_and_each_segmentation_reconcile(self):
        ws = self.workbook["Backlog_Analysis"]
        for index in range(len(ws.tables)):
            rows = self.rows("Backlog_Analysis", index)
            self.assertEqual(sum(row["Backlog_Count"] for row in rows), self.kpis["backlog"])
            self.assertAlmostEqual(sum(row["Backlog_Share"] for row in rows), 1)
        self.assertEqual(len(self.rows("Backlog_Analysis")), 6)

    def test_owner_and_actual_handler_attribution_are_separate(self):
        tickets, logs = self.data["tickets"], self.data["ticket_work_logs"]
        rows = self.rows("Agent_Performance")
        self.assertEqual(len(rows), len(self.data["agents"]))
        for row in rows:
            owned = tickets.loc[tickets.assigned_agent_id.eq(row["Agent_ID"])]
            handled = logs.loc[logs.agent_id.eq(row["Agent_ID"])]
            self.assertEqual(row["Owned_Tickets"], len(owned))
            self.assertEqual(row["Handled_Tickets"], handled.ticket_id.nunique())
            self.assertAlmostEqual(row["Handling_Hours"], handled.handling_minutes.sum() / 60)
        self.assertAlmostEqual(sum(row["Handling_Hours"] for row in rows), self.kpis["handling_hours"])
        teams = self.rows("Team_Performance")
        self.assertEqual(sum(row["Owned_Tickets"] for row in teams), len(tickets))
        self.assertAlmostEqual(sum(row["Productive_Capacity_Hours"] for row in teams), self.kpis["productive_capacity_hours"])
        self.assertNotIn("Rank", self.rows("Agent_Performance")[0])

    def test_capacity_and_insufficient_history_preserve_blanks(self):
        daily = self.rows("Workforce_Utilization")
        self.assertEqual(len(daily), len(DAYS) * self.data["agents"].team.nunique())
        for row in daily:
            capacity = row["Scheduled_Hours"] - row["Absence_Hours"] - row["Shrinkage_Hours"]
            self.assertAlmostEqual(row["Productive_Capacity_Hours"], capacity)
            self.assertLessEqual(row["Handling_Hours"], row["Physical_Attendance_Hours"] + 1e-9)
            if capacity == 0:
                self.assertIsNone(row["Utilization"])
            else:
                self.assertAlmostEqual(row["Utilization"], row["Handling_Hours"] / capacity)
        staffing = self.rows("Staffing_Analysis")
        incomplete = [row for row in staffing if row["History_Days"] < 4]
        self.assertGreater(len(incomplete), 0)
        for row in incomplete:
            self.assertIsNone(row["Forecast_Workload_Minutes"])
            self.assertIsNone(row["Required_FTE"])
            self.assertIsNone(row["Staffing_Gap"])

    def test_export_preserves_every_data_artifact(self):
        for path, digest in self.before.items():
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest, path.name)

    def test_dax_core_source_contract_matches_canonical_definitions(self):
        # A static guard against changed inputs/denominators; this does not run DAX.
        text = (ROOT / "powerbi/measures.dax").read_text(encoding="utf-8-sig")
        measures, current = {}, None
        for line in text.splitlines():
            match = re.match(r"^([A-Za-z0-9][^=]+?)\s*=\s*(.*)$", line)
            if match and not line.startswith("VAR "):
                current = match.group(1).strip()
                self.assertNotIn(current, measures)
                measures[current] = match.group(2)
            elif current and not line.lstrip().startswith("//"):
                measures[current] += line
        expected = {
            "Total Tickets": "COUNTROWS ( fact_tickets )",
            "Completed Tickets": "CALCULATE ( [Total Tickets], fact_tickets[completed] = TRUE () )",
            "Open Tickets": "CALCULATE ( [Total Tickets], fact_tickets[status] = \"open\" )",
            "Pending Tickets": "CALCULATE ( [Total Tickets], fact_tickets[status] = \"pending\" )",
            "Average CSAT": "AVERAGE ( fact_tickets[csat_score] )",
            "CSAT Response Rate %": "DIVIDE ( [CSAT Responses], [Completed Tickets] )",
            "Reopen Rate %": "DIVIDE ( [Reopened Tickets], [Total Tickets] )",
            "Backlog": "[Open Tickets] + [Pending Tickets]",
            "Backlog %": "DIVIDE ( [Backlog], [Total Tickets] )",
            "Handling Hours": "DIVIDE ( SUM ( fact_work_logs[handling_minutes] ), 60 )",
            "Scheduled Hours": "DIVIDE ( SUM ( fact_workforce_daily[scheduled_minutes] ), 60 )",
            "Absence Hours": "DIVIDE ( SUM ( fact_workforce_daily[absence_minutes] ), 60 )",
            "Shrinkage Hours": "DIVIDE ( SUM ( fact_workforce_daily[shrinkage_minutes] ), 60 )",
            "Productive Capacity Hours": "[Scheduled Hours] - [Absence Hours] - [Shrinkage Hours]",
            "Utilization %": "DIVIDE ( [Handling Hours], [Productive Capacity Hours] )",
        }
        for prefix in ["FR", "Resolution", "Overall"]:
            expected[f"{prefix} SLA Compliance %"] = f"DIVIDE ( [{prefix} SLA Met], [{prefix} SLA Met] + [{prefix} SLA Breached] )"
            expected[f"{prefix} SLA Breach %"] = f"DIVIDE ( [{prefix} SLA Breached], [{prefix} SLA Met] + [{prefix} SLA Breached] )"
            field = "fr" if prefix == "FR" else prefix.lower()
            for outcome in ["Met", "Breached", "Pending"]:
                expected[f"{prefix} SLA {outcome}"] = f'CALCULATE ( [Total Tickets], fact_tickets[{field}_sla_outcome] = "{outcome.upper()}" )'
        for name, expression in expected.items():
            self.assertEqual(re.sub(r"\s+", "", measures[name]), re.sub(r"\s+", "", expression), name)
        self.assertIn("[completed] = TRUE ()", measures["CSAT Responses"])
        self.assertIn("NOT ISBLANK", measures["CSAT Responses"])
        self.assertIn("[reopen_count] > 0", measures["Reopened Tickets"])
        queries = (ROOT / "powerbi/processed_queries.pq").read_text()
        self.assertIn('\\data\\processed\\', queries)
        self.assertNotIn(".xlsx", queries)
        self.assertNotIn("data/raw", queries)


if __name__ == "__main__":
    unittest.main()
