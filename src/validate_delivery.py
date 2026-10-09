"""Validate persisted deliverables without claiming native SQL/BI execution."""
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from urllib.parse import unquote
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from contract import ROOT, hashes, read_tables, validation_errors
from verified_metrics import summarize
from export_excel import validate_workbook

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--pipeline-log",type=Path)
    args=parser.parse_args()
    meta=json.loads((ROOT/"data/analytics/generation_metadata.json").read_text())
    assert hashes()==meta["hashes"],"Frozen raw differs"
    data=read_tables("processed")
    assert not validation_errors(data)
    k=summarize(data)
    excel=validate_workbook(ROOT/"output/customer_support_analysis.xlsx",k)
    wb=load_workbook(ROOT/"output/customer_support_analysis.xlsx",data_only=False)
    assert excel["sheet_count"]==14 and excel["charts"]==5
    assert all(c.data_type!="f" and c.data_type!="e" for s in wb for row in s for c in row)
    assert all(s.tables and s.freeze_panes for s in wb)
    assert len(wb["SLA_Analysis"].conditional_formatting)>0
    assert len(wb["Staffing_Analysis"].conditional_formatting)>0
    sources=[]
    for s in wb:
        for c in s._charts:
            assert c.series,"Empty chart series"
            for series in c.series:
                assert series.val.numRef.f,"Chart missing value range"
                sources.append(series.val.numRef.f)
    claims=json.loads((ROOT/"data/analytics/case_study_claims.json").read_text())
    cats=pd.read_csv(ROOT/"data/analytics/category_service_summary.csv")
    days=pd.read_csv(ROOT/"data/analytics/demand_day_type.csv")
    assert int(cats.tickets.sum())==k["total_tickets"]
    assert int(cats.resolution_breaches.sum())==k["resolution_sla_breached"]
    assert int(days.tickets.sum())==k["total_tickets"] and int(days.calendar_days.sum())==365
    assert np.isclose(days.iloc[0].average_daily_tickets/days.iloc[1].average_daily_tickets,k["weekday_weekend_ratio"])
    contract=pd.read_csv(ROOT/"data/analytics/kpi_contract.csv")
    assert set(contract.metric)==set(k)
    assert contract[["meaning","numerator","denominator","filters","period","unit","null_rule"]].notna().all().all()
    references=json.loads((ROOT/"powerbi/acceptance_reference.json").read_text())
    assert len(references)==41 and references[0]["tickets"]==14774
    images=[]
    for rel in ["images/workbook/workbook_overview.png","images/workbook/workbook_demand.png"]:
        p=ROOT/rel
        with Image.open(p) as im:
            im.verify()
        with Image.open(p) as im:
            size=list(im.size)
        assert min(size)>300
        images.append(dict(path=rel,size=size,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),renderer="Artifact Tool from actual XLSX; not native Excel capture"))
    public=["README.md","docs/how_to_run.md","docs/sql_analysis_guide.md","docs/data_quality_report.md",
        "docs/kpi_definitions.md","docs/05_business_insights.md","docs/portfolio_case_study.md",
        "docs/independent_validation.md","docs/excel_native_checklist.md",
        "docs/claim_verification.md","docs/data_source.md","docs/interview_questions.md","docs/delivery_report.md",
        "powerbi/README.md","powerbi/data_model.md","powerbi/dashboard_spec.md",
        "powerbi/desktop_checklist.md","sql/README.md","images/workbook/README.md"]
    broken=[]
    for rel in public:
        p=ROOT/rel
        text=p.read_text(encoding="utf-8-sig")
        assert not re.search(r"\{\{[^}]+\}\}",text),rel
        for target in re.findall(r"!?\[[^\]]*\]\(([^)]+)\)",text):
            if target.startswith(("https://","http://","#","mailto:")):
                continue
            resolved=(p.parent/unquote(target.split("#")[0])).resolve()
            if not resolved.exists() and resolved != (ROOT / "output/validation_receipt.json").resolve():
                broken.append(dict(file=rel,target=target))
    assert not broken,broken
    # Check publishable CSV fields, rather than incidental IDs or dates.
    fields=[col.lower() for frame in data.values() for col in frame.columns]
    assert not set(fields)&{"email","phone_number","customer_name","customer_email","address","message_body"}
    pipeline=None
    if args.pipeline_log:
        log=args.pipeline_log.read_text(encoding="utf-8")
        found=re.search(r"Ran (\d+) tests",log)
        assert found and "\nOK" in log and "Python workflow complete." in log
        pipeline=dict(log=str(args.pipeline_log.resolve().relative_to(ROOT)),tests=int(found.group(1)),result="OK")
    from independent_validation import build
    independent=build()
    assert independent["status"]=="PASS"
    mysql_receipt=json.loads((ROOT/"output/mysql/execution_receipt.json").read_text()) if (ROOT/"output/mysql/execution_receipt.json").exists() else None
    result=dict(independent_validation=dict(status=independent["status"],manual_samples=len(independent["manual_samples"]),
                                          bi_cohorts=len(independent["bi_cohort_checks"]),receipt="output/independent_validation.json"),
        mysql_execution_receipt=mysql_receipt,
        date_local="2026-10-08",data_type="synthetic",raw_hashes_unchanged=True,processed_rules_violations=0,
        workbook=excel,workbook_sha256=hashlib.sha256((ROOT/"output/customer_support_analysis.xlsx").read_bytes()).hexdigest(),
        portable_fallback=json.loads((ROOT/"output/fallback_validation.json").read_text()) if (ROOT/"output/fallback_validation.json").exists() else None,
        chart_value_sources=sources,kpi_contract_rows=len(contract),bi_reference_cases=len(references),
        images=images,public_links_broken=broken,pipeline=pipeline,
        native_mysql="BLOCKED before connection; actual runner exit 1; client unavailable; no server/Docker in checked paths/services; 3306 refused",
        native_excel="NOT OPENED/CAPTURED; Excel executable exists; native inventory pipe unavailable (os error 2)",
        native_powerbi="NOT BUILT/OPENED; native inventory pipe unavailable; common paths absent, WindowsApps EPERM; M/DAX not executed",
        secrets_or_personal_fields_in_processed_schema=False,pushed=False,
        runtime=dict(python=sys.version.split()[0],pandas=pd.__version__,numpy=np.__version__))
    path=ROOT/"output/validation_receipt.json"
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8",newline="\n")
    wb.close()
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
