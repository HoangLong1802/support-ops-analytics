"""Export verified processed-data analysis as a readable Excel workbook."""
from pathlib import Path
import json
import math
import os
import shutil
import subprocess
import tempfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from contract import ROOT, DAYS, SNAPSHOT, ZONE, hashes, read_tables, validation_errors
from verified_metrics import ratio, summarize, ticket_metrics, workforce_metrics
from staffing_analysis import estimate

OUTPUT = ROOT / "output/customer_support_analysis.xlsx"
SHEETS = ["README", "Executive_KPIs", "Demand_Analysis", "SLA_Analysis",
          "Category_Analysis", "CSAT_Reopen", "Backlog_Analysis", "Agent_Performance",
          "Team_Performance", "Workforce_Utilization", "Staffing_Analysis",
          "Data_Quality", "Cleaning_Summary", "Quarantine_Summary"]
KPI_LABELS = {
    "total_tickets": "Total Tickets", "completed_tickets": "Completed Tickets",
    "open_tickets": "Open Tickets", "pending_tickets": "Pending Tickets",
    "average_first_response_minutes": "Average First Response Minutes",
    "average_resolution_hours": "Average Resolution Hours", "median_resolution_hours": "Median Resolution Hours",
    "p90_resolution_hours": "P90 Resolution Hours", "p95_resolution_hours": "P95 Resolution Hours",
    "average_csat": "Average CSAT", "csat_responses": "CSAT Responses",
    "csat_response_rate": "CSAT Response Rate %", "reopened_tickets": "Reopened Tickets",
    "reopen_rate": "Reopen Rate %", "backlog": "Backlog Count", "backlog_rate": "Backlog %",
    "weekday_weekend_ratio": "Weekday / Weekend Average Arrivals",
    "handling_hours": "Handling Hours", "productive_capacity_hours": "Productive Capacity Hours",
    "utilization": "Utilization %", "median_team_day_utilization": "Median Team-Day Utilization %",
}
for prefix, label in [("fr", "First Response"), ("resolution", "Resolution"), ("overall", "Overall")]:
    for outcome in ["met", "breached", "pending", "eligible"]:
        KPI_LABELS[f"{prefix}_sla_{outcome}"] = f"{label} SLA {outcome.title()}"
    KPI_LABELS[f"{prefix}_sla_compliance"] = f"{label} SLA Compliance %"
    KPI_LABELS[f"{prefix}_breach_rate"] = f"{label} SLA Breach %"
for hours in [24, 48, 72]:
    KPI_LABELS[f"backlog_over_{hours}_hours"] = f"Backlog >{hours} Hours"


def verified_inputs(data):
    errors = validation_errors(data)
    if errors:
        raise ValueError("Invalid processed source: " + "; ".join(errors))
    metadata = json.loads((ROOT / "data/analytics/generation_metadata.json").read_text())
    if hashes() != metadata["hashes"]:
        raise ValueError("Frozen raw checksum differs")
    kpis = summarize(data)
    saved = pd.read_csv(ROOT / "data/analytics/verified_kpis.csv").set_index("metric").value
    for key, value in kpis.items():
        if not np.isclose(value, saved[key], rtol=1e-12, atol=1e-9, equal_nan=True):
            raise ValueError(f"Saved KPI differs: {key}; refresh verified_metrics.py first")
    return kpis


def executive_rows(kpis):
    definitions = pd.read_csv(ROOT / "data/analytics/verified_kpis.csv").set_index("metric").definition.to_dict()
    for prefix in ["fr", "resolution", "overall"]:
        definitions[f"{prefix}_breach_rate"] = "BREACHED / (MET + BREACHED); PENDING excluded"
    definitions.update({
        "total_tickets": "One retained ticket snapshot per ticket_id",
        "completed_tickets": "Tickets with status resolved or closed",
        "open_tickets": "Tickets with status open at the snapshot",
        "pending_tickets": "Tickets with status pending at the snapshot",
        "average_first_response_minutes": "Mean observed first-response elapsed minutes; missing events excluded",
        "average_resolution_hours": "Mean completed-ticket elapsed hours; includes waiting, not handling effort",
        "average_csat": "Mean valid completed-ticket survey score (1–5); no-response excluded",
        "backlog": "Open + pending tickets at the fixed snapshot",
        "backlog_rate": "Snapshot backlog / all retained tickets",
        "handling_hours": "Actual retained work-log minutes / 60; attributed to handlers",
        "productive_capacity_hours": "Sum(scheduled − absence − shrinkage) / 60",
    })
    return pd.DataFrame([{"Metric": KPI_LABELS[key], "Value": value, "Definition": definitions[key]}
                         for key, value in kpis.items()])


def ticket_cohorts(tickets, columns):
    rows = []
    for keys, group in tickets.groupby(columns, dropna=False, sort=True):
        keys = keys if isinstance(keys, tuple) else (keys,)
        row = dict(zip(columns, keys))
        row.update(Ticket_Count=len(group), Completed_Tickets=int(group.completed.sum()),
                   Ticket_Share=ratio(len(group), len(tickets)),
                   CSAT_Response_Count=int(group.csat_score.count()),
                   CSAT_Response_Rate=ratio(group.csat_score.count(), group.completed.sum()),
                   Average_CSAT=group.csat_score.mean(),
                   Average_First_Response_Minutes=group.first_response_minutes.mean(),
                   Average_Resolution_Hours=group.resolution_minutes.mean() / 60,
                   Reopen_Rate=group.reopen_count.gt(0).mean(),
                   Technical_Case_Share=group.category.eq("technical_support").mean(),
                   High_Urgent_Share=group.priority.isin(["high", "urgent"]).mean())
        for prefix, label in [("fr", "FR"), ("resolution", "Resolution"), ("overall", "Overall")]:
            outcomes = group[f"{prefix}_sla_outcome"]
            met, breached = int(outcomes.eq("MET").sum()), int(outcomes.eq("BREACHED").sum())
            row[f"{label}_SLA_Eligible"] = met + breached
            row[f"{label}_SLA_Compliance"] = ratio(met, met + breached)
            row[f"{label}_SLA_Breach_Rate"] = ratio(breached, met + breached)
        rows.append(row)
    return pd.DataFrame(rows)


def demand_sections(tickets):
    daily = tickets.groupby("local_created_date").size().reindex(DAYS, fill_value=0)
    monthly = daily.groupby(daily.index.to_period("M")).sum()
    trend = pd.DataFrame({"Period": monthly.index.to_timestamp().date, "Ticket_Count": monthly.values})
    trend["Month"] = pd.to_datetime(trend.Period).dt.strftime("%b%y")
    trend["Previous_Period"] = trend.Period.shift(1)
    trend["Previous_Ticket_Count"] = trend.Ticket_Count.shift(1)
    trend["Change"] = trend.Ticket_Count - trend.Previous_Ticket_Count
    trend["Change_Percent"] = trend.Change / trend.Previous_Ticket_Count.replace(0, np.nan)
    weekday = pd.DataFrame({"Weekday_Number": range(7)})
    # Calendar coverage, rather than only dates with arrivals, supplies the denominator.
    weekday["Segment"] = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    weekday["Ticket_Count"] = [int(daily[DAYS.dayofweek == day].sum()) for day in range(7)]
    weekday["Calendar_Days"] = [int((DAYS.dayofweek == day).sum()) for day in range(7)]
    weekday["Average_Daily_Tickets"] = weekday.Ticket_Count / weekday.Calendar_Days
    weekday["Ticket_Share"] = weekday.Ticket_Count / len(tickets)
    hours = tickets.groupby("local_created_hour").size().reindex(range(24), fill_value=0)
    hourly = pd.DataFrame({"Local_Hour": hours.index, "Ticket_Count": hours.values, "Ticket_Share": hours.values / len(tickets)})
    sections = [("Monthly ticket trend", trend), ("Weekday demand", weekday), ("Local-hour arrivals", hourly)]
    for dimension in ["channel", "priority"]:
        count = tickets[dimension].value_counts().sort_index()
        sections.append((dimension.title() + " mix", pd.DataFrame({"Dimension": dimension, "Segment": count.index,
                         "Ticket_Count": count.values, "Ticket_Share": count.values / len(tickets)})))
    return sections


def sla_rows(tickets):
    rows = []
    dimensions = [("Overall", None), ("Channel", "channel"), ("Priority", "priority"),
                  ("Category", "category"), ("Final_Owner_Team", "owner_team")]
    for dimension, column in dimensions:
        cohorts = [("All retained tickets", tickets)] if column is None else tickets.groupby(column, dropna=False)
        for segment, group in cohorts:
            for prefix, label in [("fr", "First Response"), ("resolution", "Resolution"), ("overall", "Overall")]:
                outcome = group[f"{prefix}_sla_outcome"]
                met, breached, pending = [int(outcome.eq(value).sum()) for value in ["MET", "BREACHED", "PENDING"]]
                rows.append({"SLA_Type": label, "Dimension": dimension, "Segment": segment,
                             "Met": met, "Breached": breached, "Pending": pending, "Eligible": met + breached,
                             "Compliance_Rate": ratio(met, met + breached), "Breach_Rate": ratio(breached, met + breached)})
    return pd.DataFrame(rows)


def backlog_sections(tickets):
    backlog = tickets.loc[~tickets.completed].copy()
    backlog["Age_Hours"] = backlog.backlog_age_minutes / 60
    labels = ["<24h", "24–48h", "48–72h", "3–7 days", "7–30 days", "30+ days"]
    backlog["age_bucket"] = pd.cut(backlog.Age_Hours, [0, 24, 48, 72, 168, 720, np.inf], labels=labels, right=False)
    sections = []
    for dimension in ["age_bucket", "status", "category", "priority", "owner_team"]:
        rows = []
        for segment, group in backlog.groupby(dimension, observed=False, dropna=False):
            rows.append({"Dimension": dimension, "Segment": str(segment), "Backlog_Count": len(group),
                         "Backlog_Share": ratio(len(group), len(backlog)), "Median_Age_Hours": group.Age_Hours.median(),
                         "P90_Age_Hours": group.Age_Hours.quantile(.9)})
        sections.append((dimension.replace("_", " ").title(), pd.DataFrame(rows)))
    return sections


def performance_rows(data, tickets, workforce, by_team=False):
    owner_column = "owner_team" if by_team else "assigned_agent_id"
    output_key = "Team" if by_team else "Agent_ID"
    owners = ticket_cohorts(tickets, [owner_column]).rename(columns={owner_column: output_key,
                    "Ticket_Count": "Owned_Tickets", "CSAT_Response_Count": "CSAT_Responses"})
    logs = data["ticket_work_logs"].merge(data["agents"][["agent_id", "team"]], on="agent_id", validate="many_to_one")
    key = "team" if by_team else "agent_id"
    effort = logs.groupby(key).agg(Handled_Tickets=("ticket_id", "nunique"), Handling_Minutes=("handling_minutes", "sum"))
    capacity = workforce.groupby(key).productive_minutes.sum().rename("Productive_Minutes")
    effort = effort.join(capacity, how="outer").reset_index().rename(columns={key: output_key})
    if by_team:
        base = pd.DataFrame({"Team": sorted(set(data["agents"].team) | set(tickets.owner_team))})
    else:
        base = data["agents"].rename(columns={"agent_id": "Agent_ID", "team": "Team", "hire_date": "Hire_Date"})
    out = base.merge(owners, on=output_key, how="left", validate="one_to_one").merge(effort, on=output_key, how="left", validate="one_to_one")
    for column in ["Owned_Tickets", "Completed_Tickets", "CSAT_Responses", "Handled_Tickets", "Handling_Minutes", "Productive_Minutes"]:
        out[column] = out[column].fillna(0)
    out["Handling_Hours"] = out.pop("Handling_Minutes") / 60
    out["Productive_Capacity_Hours"] = out.pop("Productive_Minutes") / 60
    out["Utilization"] = out.Handling_Hours / out.Productive_Capacity_Hours.replace(0, np.nan)
    if not by_team:
        out["Hire_Date"] = pd.to_datetime(out.Hire_Date).dt.date
    context = ["Team"] if by_team else ["Agent_ID", "Team", "Hire_Date"]
    columns = context + ["Owned_Tickets", "Handled_Tickets", "Handling_Hours", "Productive_Capacity_Hours", "Utilization",
              "Completed_Tickets", "Average_Resolution_Hours", "FR_SLA_Eligible", "FR_SLA_Compliance",
              "Resolution_SLA_Eligible", "Resolution_SLA_Compliance", "Overall_SLA_Eligible", "Overall_SLA_Compliance",
              "Average_CSAT", "CSAT_Responses", "CSAT_Response_Rate", "Reopen_Rate", "Technical_Case_Share", "High_Urgent_Share"]
    return out[columns]


def quality_tables(tickets):
    assessment = json.loads((ROOT / "data/analytics/quality_audit.json").read_text())
    cleaning = json.loads((ROOT / "data/analytics/cleaning_audit.json").read_text())
    quality = []
    for issue in assessment["issues"]:
        denominator = cleaning["raw"][issue["dataset"]]
        quality.append({"Issue_Type": "Rule issue", "Severity": issue["severity"], "Dataset": issue["dataset"],
                        "Issue": issue["issue"], "Affected_Rows": issue["count"], "Denominator": denominator,
                        "Percentage": ratio(issue["count"], denominator), "Business_Impact": issue["business_impact"],
                        "Recommended_Action": issue["treatment"]})
    long_count = int(tickets.resolution_minutes.gt(200 * 24 * 60).sum())
    quality.append({"Issue_Type": "Statistical observation", "Severity": "ANOMALY", "Dataset": "processed tickets",
                    "Issue": "Completed resolution exceeds 200 days; logically valid", "Affected_Rows": long_count,
                    "Denominator": int(tickets.completed.sum()), "Percentage": ratio(long_count, tickets.completed.sum()),
                    "Business_Impact": "Long waits influence the mean; they are not automatically data errors.",
                    "Recommended_Action": "KEEP; review waiting context and report tail percentiles"})
    decisions = []
    def decision(issue, treatment, count, reason, impact):
        decisions.append({"Issue": issue, "Treatment": treatment, "Rows_Affected": count,
                          "Reason": reason, "Analytical_Impact": impact})
    for dataset, count in cleaning["exact_copies_removed"].items():
        if count:
            decision(dataset + ": exact copies", "REMOVE", count, "Every source field repeats the observation", "Avoid volume/effort inflation")
    for key, treatment, reason, impact in [
        ("channel_normalized", "NORMALIZE", "Case and whitespace preserve meaning", "Restore channel groups"),
        ("category_restored", "DERIVE", "Known subcategory has one approved parent", "Retain supported case mix"),
        ("invalid_csat_set_null", "NULL", "Cannot recover or clamp an invalid score", "Retain ticket; exclude invalid survey value")]:
        decision(key, treatment, cleaning["transformations"][key], reason, impact)
    for dataset, reasons in cleaning["reason_counts"].items():
        for reason, count in reasons.items():
            evidence = {
                "conflicting_key": "No source precedence identifies the correct conflicting version",
                "missing_hierarchy": "A category cannot identify its missing subcategory",
                "invalid_hierarchy": "Neither conflicting field identifies the correct pair",
                "completed_missing_fields": "Completion status cannot supply missing event or owner evidence",
                "response_order": "Response precedes creation or falls beyond the snapshot",
                "unknown_owner": "The agent reference has no valid source record",
                "unknown_ticket": "Ticket parent is missing, ambiguous or excluded",
                "missing_workforce": "No retained capacity record for the handler/day",
                "invalid_handling": "Handling must be a positive integer number of minutes",
                "invalid_capacity": "Absence plus shrinkage exceeds scheduled minutes",
            }
            decision(dataset + ": " + reason, "QUARANTINE", count, evidence[reason], "Preserve evidence; exclude unreliable analytical facts")
    decision("Resolution exceeds 200 days; logically valid", "KEEP", long_count, "No duration upper bound or supported replacement", "Preserve genuine elapsed-time tails")
    reconciliation = pd.DataFrame([{"Dataset": name, "Raw_Rows": cleaning["raw"][name],
                      "Exact_Copies_Removed": cleaning["exact_copies_removed"][name], "Quarantined_Rows": cleaning["quarantined"][name],
                      "Processed_Rows": cleaning["processed"][name]} for name in cleaning["raw"]])
    quarantine = pd.DataFrame([{"Dataset": name, "Reason": reason, "Row_Count": count,
                  "Percentage_Of_Source": ratio(count, cleaning["raw"][name])}
                  for name, reasons in cleaning["reason_counts"].items() for reason, count in reasons.items()])
    return pd.DataFrame(quality), pd.DataFrame(decisions), reconciliation, quarantine


def workbook_sections(data, kpis):
    tickets = ticket_metrics(data["tickets"], data["sla_policies"])
    tickets["owner_team"] = tickets.assigned_agent_id.map(data["agents"].set_index("agent_id").team).fillna("unassigned")
    tickets["Reopen_Status"] = np.where(tickets.reopen_count.gt(0), "reopened", "no_recorded_reopen")
    workforce = workforce_metrics(data)
    staffing = estimate(data)
    saved_staffing = pd.read_csv(ROOT / "data/analytics/staffing_daily.csv")
    comparison = staffing.copy()
    comparison["work_date"] = comparison.work_date.dt.strftime("%Y-%m-%d")
    pd.testing.assert_frame_equal(comparison.reset_index(drop=True), saved_staffing, check_dtype=False, rtol=1e-12, atol=1e-9)
    daily = workforce.groupby(["work_date", "team"])[["scheduled_minutes", "absence_minutes", "shrinkage_minutes",
              "productive_minutes", "physical_attendance_minutes", "handling_minutes"]].sum()
    calendar = pd.MultiIndex.from_product([DAYS.strftime("%Y-%m-%d"), sorted(data["agents"].team.unique())], names=["work_date", "team"])
    daily = daily.reindex(calendar, fill_value=0).reset_index()
    daily = daily.rename(columns={"work_date": "Date", "team": "Team"})
    for source, label in [("scheduled_minutes", "Scheduled_Hours"), ("absence_minutes", "Absence_Hours"),
                          ("shrinkage_minutes", "Shrinkage_Hours"), ("productive_minutes", "Productive_Capacity_Hours"),
                          ("physical_attendance_minutes", "Physical_Attendance_Hours"), ("handling_minutes", "Handling_Hours")]:
        daily[label] = daily.pop(source) / 60
    daily["Utilization"] = daily.Handling_Hours / daily.Productive_Capacity_Hours.replace(0, np.nan)
    daily["Date"] = pd.to_datetime(daily.Date).dt.date
    staffing = staffing.rename(columns={"work_date": "Date", "team": "Team", "forecast_handling_minutes": "Forecast_Workload_Minutes",
              "forecast_history_days": "History_Days", "productive_minutes_per_fte": "Productive_Minutes_Per_FTE",
              "target_utilization": "Target_Utilization", "required_fte": "Required_FTE", "available_fte": "Available_FTE",
              "staffing_gap_fte": "Staffing_Gap"})
    staffing["Date"] = staffing.Date.dt.date
    staffing = staffing[["Date", "Team", "History_Days", "Forecast_Workload_Minutes", "Productive_Minutes_Per_FTE",
                         "Target_Utilization", "Required_FTE", "Available_FTE", "Staffing_Gap"]]
    scores = tickets.csat_score.value_counts().reindex(range(1, 6), fill_value=0).sort_index()
    csat_sections = [("Valid CSAT distribution", pd.DataFrame({"CSAT_Score": scores.index, "Response_Count": scores.values,
                     "Response_Share": scores.values / tickets.csat_score.count()}))]
    for column in ["channel", "category", "customer_type", "Reopen_Status"]:
        cohort = ticket_cohorts(tickets, [column])
        cohort["SLA_Breach_Rate"] = cohort.Overall_SLA_Breach_Rate
        csat_sections.append((column.replace("_", " ").title(), cohort[[column, "Ticket_Count", "Completed_Tickets",
                             "CSAT_Response_Count", "CSAT_Response_Rate", "Average_CSAT", "Average_Resolution_Hours",
                             "SLA_Breach_Rate", "Overall_SLA_Eligible", "Reopen_Rate"]]))
    reopen = tickets.reopen_count.value_counts().sort_index()
    csat_sections.append(("Recorded reopen count distribution", pd.DataFrame({"Reopen_Count": reopen.index,
                          "Ticket_Count": reopen.values, "Ticket_Share": reopen.values / len(tickets)})))
    quality, cleaning, reconciliation, quarantine = quality_tables(tickets)
    categories = ticket_cohorts(tickets, ["category", "subcategory"]).rename(columns={"category": "Category", "subcategory": "Subcategory"})
    categories = categories[["Category", "Subcategory", "Ticket_Count", "Ticket_Share", "Completed_Tickets",
                  "FR_SLA_Eligible", "FR_SLA_Breach_Rate", "Resolution_SLA_Eligible", "Resolution_SLA_Breach_Rate",
                  "Overall_SLA_Eligible", "Overall_SLA_Breach_Rate", "Average_First_Response_Minutes",
                  "Average_Resolution_Hours", "Average_CSAT", "CSAT_Response_Count", "CSAT_Response_Rate", "Reopen_Rate"]]
    descriptions = {
        "README": "Scope, caveats, sheet navigation and refresh command",
        "Executive_KPIs": "Verified snapshot KPIs and their definitions",
        "Demand_Analysis": "Monthly, weekday, hourly, channel and priority arrival comparisons",
        "SLA_Analysis": "MET/BREACHED/PENDING outcomes and eligible denominators by cohort",
        "Category_Analysis": "Category/subcategory service, CSAT, reopen and sample context",
        "CSAT_Reopen": "Survey distribution, participation and noncausal reopen comparisons",
        "Backlog_Analysis": "Fixed-snapshot status, age buckets and ownership segments",
        "Agent_Performance": "Final-owner outcomes beside actual-handler effort; no ranking",
        "Team_Performance": "Owner cohorts, actual team handling and productive capacity",
        "Workforce_Utilization": "Calendar team-day effort, schedule, absence, shrinkage and capacity",
        "Staffing_Analysis": "Retrospective daily equivalent-FTE planning; no exact scheduling",
        "Data_Quality": "Detected source issues separated from valid long-duration observations",
        "Cleaning_Summary": "Implemented treatments and complete source-row reconciliation",
        "Quarantine_Summary": "Overlapping exclusion reasons; source denominators, no raw dump",
    }
    intro = pd.DataFrame([
        ("Project", "Customer Support Operations Analytics"), ("Dataset", "Dữ liệu vận hành support mô phỏng"),
        ("Workflow", "Raw → Quality → Cleaning → Analysis → Excel → Power BI → Insights"),
        ("Reporting timezone", ZONE), ("Snapshot", str(SNAPSHOT.tz_convert(ZONE))),
        ("Analytical source", "Nguồn: data/processed/ và kết quả data/analytics/. Excel là sản phẩm trình bày; Power BI đọc processed CSV."),
        ("Limitations", "Dữ liệu mô phỏng; Reopen Rate không phải FCR; backlog chỉ tại snapshot; elapsed resolution không phải effort; capacity theo ngày; staffing là ước tính; không kết luận nhân quả."),
        ("Units / blanks", "Rates lưu dạng fraction 0–1, hiển thị %. Null hoặc mẫu số 0 để trống. Các giá trị là kết quả Python đã export, không phải công thức Excel hoặc PivotTable."),
        ("Attribution", "Ticket outcomes: final owner. Effort: actual handler. Handled_Tickets is distinct within each agent/team and is not additive across handlers."),
        ("Backlog boundaries", "Age buckets are half-open: 24–48h means 24 <= age < 48. KPI >48h uses strictly greater than 48, as defined in the project."),
        ("Refresh", "python src/export_excel.py"),
    ], columns=["Field", "Value"])
    return {
        "README": [("KPI Overview / Tổng quan", pd.DataFrame([
            ("Total Tickets", kpis["total_tickets"], "ticket; unique ID"),
            ("Resolution SLA Compliance %", kpis["resolution_sla_compliance"], "MET / eligible"),
            ("FR SLA Compliance %", kpis["fr_sla_compliance"], "MET / eligible"),
            ("Backlog Count", kpis["backlog"], "open + pending"),
            ("Backlog >48 Hours", kpis["backlog_over_48_hours"], "ticket; age >48h"),
            ("Average CSAT", kpis["average_csat"], "score /5; respondents"),
            ("CSAT Response Rate %", kpis["csat_response_rate"], "valid CSAT / completed"),
            ("Weekday / Weekend Average Arrivals", kpis["weekday_weekend_ratio"], "ratio; calendar days")
            ], columns=["Metric", "Value", "Unit / Scope"])),
            ("SLA compliance", pd.DataFrame([("First Response", kpis["fr_sla_compliance"]),
                ("Resolution", kpis["resolution_sla_compliance"]), ("Overall", kpis["overall_sla_compliance"])], columns=["Stage", "Compliance_Rate"])),
            ("Workbook guide", pd.DataFrame(list(descriptions.items()), columns=["Sheet", "Contents"])),
            ("Project scope / Phạm vi", intro)],
        "Executive_KPIs": [("Verified KPI values", executive_rows(kpis))],
        "Demand_Analysis": demand_sections(tickets), "SLA_Analysis": [("Snapshot SLA by cohort", sla_rows(tickets))],
        "Category_Analysis": [("Category and subcategory", categories)], "CSAT_Reopen": csat_sections,
        "Backlog_Analysis": backlog_sections(tickets),
        "Agent_Performance": [("Owners and handlers", performance_rows(data, tickets, workforce))],
        "Team_Performance": [("Ownership and actual team effort", performance_rows(data, tickets, workforce, by_team=True))],
        "Workforce_Utilization": [("Calendar team-day workload and capacity", daily)],
        "Staffing_Analysis": [("Daily planning estimate", staffing)],
        "Data_Quality": [("Rule findings and statistical observations", quality)],
        "Cleaning_Summary": [("Implemented treatments", cleaning), ("Source-row reconciliation", reconciliation)],
        "Quarantine_Summary": [("Reason counts (can overlap)", quarantine)],
    }, descriptions


def number_format(column):
    if column == "Date" or column == "Hire_Date":
        return "yyyy-mm-dd"
    if column in {"Period", "Previous_Period"}:
        return "yyyy-mm"
    if any(word in column for word in ["Rate", "Share", "Compliance", "Percentage", "Utilization", "Percent"]):
        return "0.00%"
    if any(word in column for word in ["Hours", "Minutes", "FTE", "Staffing_Gap", "Average_CSAT", "Average_Daily"]):
        return "#,##0.00"
    return "#,##0"


def legacy_write_workbook(path, sections, descriptions):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    locations = {}
    with pd.ExcelWriter(path, engine="openpyxl", date_format="yyyy-mm-dd", datetime_format="yyyy-mm-dd") as writer:
        for sheet, tables in sections.items():
            row = 2
            for index, (title, frame) in enumerate(tables):
                numeric = frame.select_dtypes(include="number")
                if np.isinf(numeric.to_numpy(dtype=float)).any():
                    raise ValueError(f"Infinite analytical value: {sheet}/{title}")
                frame.to_excel(writer, sheet_name=sheet, startrow=row + 1, index=False, na_rep="")
                ws = writer.sheets[sheet]
                ws.cell(row=row + 1, column=1, value=title).font = Font(name="Calibri", bold=True, size=11)
                header = row + 2
                table = Table(displayName=f"{sheet}_T{index + 1}", ref=f"A{header}:{ws.cell(header + len(frame), len(frame.columns)).coordinate}")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(table)
                for cell in ws[header]:
                    cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24445C")
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[header].height = 34
                for column, name in enumerate(frame.columns, start=1):
                    width = 70 if name in {"Definition", "Value", "Contents", "Business_Impact", "Reason", "Analytical_Impact", "Recommended_Action"} else min(30, max(16, len(name) + 2))
                    ws.column_dimensions[ws.cell(header, column).column_letter].width = max(width, ws.column_dimensions[ws.cell(header, column).column_letter].width)
                    for cells in ws.iter_rows(min_row=header + 1, max_row=header + len(frame), min_col=column, max_col=column):
                        cell = cells[0]
                        cell.font = Font(name="Calibri", size=11)
                        cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
                        cell.number_format = number_format(name)
                for excel_row in range(header + 1, header + len(frame) + 1):
                    lines = max((math.ceil(len(str(cell.value)) / max(10, ws.column_dimensions[cell.column_letter].width - 2))
                                 for cell in ws[excel_row] if isinstance(cell.value, str)), default=1)
                    ws.row_dimensions[excel_row].height = min(360, max(18, lines * 15))
                locations[(sheet, title)] = (header, header + len(frame), list(frame.columns))
                if index == 0:
                    ws.freeze_panes = f"A{header + 1}"
                row = header + len(frame) + 2
                if sheet == "README" and index == 0:
                    row = 22
            ws.sheet_view.showGridLines = False
            ws.cell(1, 1, "Scope").font = Font(name="Calibri", bold=True)
            scope = descriptions[sheet]
            if sheet in {"SLA_Analysis", "Category_Analysis", "CSAT_Reopen", "Agent_Performance", "Team_Performance"}:
                scope += "; SLA denominator = MET + BREACHED; owner outcomes differ from handler effort."
            if sheet in {"Data_Quality", "Cleaning_Summary", "Quarantine_Summary"}:
                scope += "; reason counts can overlap; distinct totals reconcile separately."
            ws.cell(1, 2, scope).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 58
        ws = writer.sheets["Executive_KPIs"]
        header, last, _ = locations[("Executive_KPIs", "Verified KPI values")]
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 38, 22, 80
        for row in range(header + 1, last + 1):
            metric = ws.cell(row, 1).value
            decimal = not metric.startswith("Backlog") and any(word in metric for word in ["Hours", "Minutes", "Average CSAT", "Average Arrivals"])
            ws.cell(row, 2).number_format = "0.00%" if "%" in metric else ("#,##0.00" if decimal else "#,##0")
        guide = writer.sheets["README"]
        first, last, _ = locations[("README", "Workbook guide")]
        for row in range(first + 1, last + 1):
            cell = guide.cell(row, 1)
            cell.hyperlink = f"#'{cell.value}'!A1"
            cell.font = Font(name="Calibri", color="0563C1", underline="single")
        charts = [("README", "SLA compliance", "SLA compliance tại snapshot", "Stage", "Compliance_Rate", "F4", False),
                  ("Demand_Analysis", "Monthly ticket trend", "Monthly arrivals", "Month", "Ticket_Count", "I3", True),
                  ("SLA_Analysis", "Snapshot SLA by cohort", "Overall-cohort SLA compliance", "SLA_Type", "Compliance_Rate", "M3", False),
                  ("Backlog_Analysis", "Age Bucket", "Snapshot backlog by age", "Segment", "Backlog_Count", "H3", False),
                  ("Team_Performance", "Ownership and actual team effort", "Observed team utilization", "Team", "Utilization", "V3", False)]
        for sheet, title, chart_title, category, value, anchor, line in charts:
            ws = writer.sheets[sheet]
            first, last, columns = locations[(sheet, title)]
            if sheet == "SLA_Analysis":
                last = first + 3
            chart = LineChart() if line else BarChart()
            chart.title, chart.width, chart.height = chart_title, 18, 9
            chart.add_data(Reference(ws, min_col=columns.index(value) + 1, min_row=first, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=columns.index(category) + 1, min_row=first + 1, max_row=last))
            chart.legend = None
            chart.y_axis.title = "Rate" if value in {"Utilization", "Compliance_Rate"} else "Tickets"
            if value in {"Utilization", "Compliance_Rate"}:
                chart.y_axis.numFmt = "0%"
            ws.add_chart(chart, anchor)
        for sheet, ref in [("SLA_Analysis", "I5:I55"), ("Staffing_Analysis", "I5:I1099")]:
            writer.sheets[sheet].conditional_formatting.add(ref, ColorScaleRule(start_type="min", start_color="EEF4F7", end_type="max", end_color="C54D4D"))
        overview = writer.sheets["README"]
        for column, width in [("A",35),("B",24),("C",38)]:
            overview.column_dimensions[column].width = width
        for row in overview.iter_rows():
            lines = max((math.ceil(len(str(c.value)) / max(10, overview.column_dimensions[c.column_letter].width - 2)) for c in row if isinstance(c.value,str)), default=1)
            overview.row_dimensions[row[0].row].height = max(25,min(360,lines*17))
        for row in range(5,13):
            label = overview.cell(row,1).value
            overview.cell(row,2).number_format = "0.00%" if "%" in label else ("#,##0.00" if any(x in label for x in ["Average CSAT","Weekday"]) else "#,##0")



def write_workbook(path, sections, descriptions):
    """Use Artifact Tool when available; retain portable openpyxl fallback."""
    node = os.environ.get("NODE_EXE") or shutil.which("node")
    try:
        available = node and subprocess.run([node, "--input-type=module", "-e", "await import('@oai/artifact-tool')"], cwd=ROOT, capture_output=True).returncode == 0
    except OSError:
        available = False
    if not available:
        return legacy_write_workbook(path, sections, descriptions)
    payload = {"sections": {name: [{"title": title, "columns": list(frame.columns),
                  "rows": json.loads(frame.to_json(orient="values", date_format="iso", double_precision=15))}
                  for title, frame in tables] for name, tables in sections.items()}}
    with tempfile.TemporaryDirectory() as temp:
        spec_path = Path(temp) / "workbook.json"
        spec_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        args = [node, str(ROOT / "src/workbook_builder.mjs"), str(spec_path), str(Path(path).resolve())]
        if Path(path).resolve() == OUTPUT.resolve():
            args += [str(ROOT / "images/workbook")]
        result = subprocess.run(args, cwd=ROOT, capture_output=True, text=True, encoding="utf-8")
        if result.returncode:
            raise RuntimeError(result.stdout + result.stderr)
        print(result.stdout)

def validate_workbook(path, kpis):
    workbook = load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != SHEETS:
            raise ValueError("Unexpected sheet order or missing sheets")
        for ws in workbook:
            if ws.max_row < 5 or not ws.tables or not ws.freeze_panes:
                raise ValueError(f"Empty or unformatted sheet: {ws.title}")
            for cells in ws.iter_rows():
                for cell in cells:
                    if cell.data_type == "f":
                        raise ValueError("Workbook must contain explicit analytical results, not hidden formulas")
                    if isinstance(cell.value, float) and not math.isfinite(cell.value):
                        raise ValueError("Nonfinite Excel value")
                    if isinstance(cell.value, str) and cell.value.strip().lower() in {"nan", "inf", "infinity", "-inf", "-infinity"}:
                        raise ValueError("Nonfinite value exported as text")
        ws = workbook["Executive_KPIs"]
        observed = {ws.cell(row, 1).value: ws.cell(row, 2) for row in range(5, ws.max_row + 1)}
        for key, expected in kpis.items():
            cell = observed[KPI_LABELS[key]]
            if pd.isna(expected):
                valid = cell.value is None
            else:
                valid = isinstance(cell.value, (int, float)) and np.isclose(cell.value, expected, rtol=1e-12, atol=1e-9)
            if not valid:
                raise ValueError(f"Excel KPI differs: {key}")
            if "%" in KPI_LABELS[key] and cell.number_format != "0.00%":
                raise ValueError(f"Incorrect percentage format: {key}")
        return {"validation": "PASS", "sheet_count": len(workbook.sheetnames), "kpis_matched": len(kpis),
                "numeric_tolerance": "rtol=1e-12; atol=1e-9; displayed rates rounded to 2 decimals", "charts": sum(len(ws._charts) for ws in workbook)}
    finally:
        workbook.close()


def export(path=OUTPUT):
    before = hashes()
    data = read_tables("processed")
    kpis = verified_inputs(data)
    sections, descriptions = workbook_sections(data, kpis)
    write_workbook(path, sections, descriptions)
    result = validate_workbook(path, kpis)
    if hashes() != before:
        raise ValueError("Excel export modified raw data")
    return result


if __name__ == "__main__":
    result = export()
    print(json.dumps({"workbook": str(OUTPUT), **result}, indent=2))
